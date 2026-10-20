"""pH, CaCO3 e tessitura 0-30 cm per ogni delineazione della Carta dei suoli 1:50.000.

Il pacchetto open data della carta (Carta_Suoli_50k.zip) contiene, oltre ai
poligoni, un foglio Excel con i suoli di ogni delineazione (percentuali e sito
rappresentativo) e le analisi di laboratorio per orizzonte dei siti.

Metodo (scelta di progetto, documentata in FASE0.md):
- per ogni delineazione si prende il suolo prevalente (percentuale più alta);
- del suo sito rappresentativo si usano gli orizzonti che cadono in 0-30 cm,
  con media pesata sullo spessore dentro lo strato;
- la classe tessiturale USDA si calcola da sabbia/limo/argilla con le regole
  del triangolo USDA; la traduzione nelle categorie del modello passa da
  lookups/suoli_tessitura.csv.
Nessun valore stimato: dove manca il dato, resta vuoto.

I nomi di fogli e colonne stanno in sources.toml ([sources.suoli.prepare]).
"""

from __future__ import annotations

import math
from collections import defaultdict
from pathlib import Path

from ..config import Source
from ..errors import ConfigError, NotConfigured
from ..lookups import norm_key

OUT_LAYER = "suoli_0_30"
OUT_NAME = "suoli_0_30.gpkg"

REQUIRED = [
    "shp_join",
    "soils_sheet",
    "soils_join",
    "soil_code",
    "soil_name",
    "soil_perc",
    "soil_site",
    "sites_sheet",
    "site_id",
    "top",
    "bottom",
    "sand",
    "silt",
    "clay",
    "caco3",
    "ph",
]


# --- funzioni pure ------------------------------------------------------------


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        f = float(str(v).replace(",", "."))
    except ValueError:
        return None
    return None if math.isnan(f) else f


def usda_texture(sand: float | None, silt: float | None, clay: float | None) -> str | None:
    """Classe tessiturale USDA (nome inglese, come il campo `usda` del modello).

    Regole del triangolo USDA (Soil Survey Manual). Le tre frazioni si
    normalizzano a 100. None se manca un valore.
    """
    vals = [sand, silt, clay]
    if any(v is None or (isinstance(v, float) and math.isnan(v)) for v in vals):
        return None
    total = sum(vals)
    if total <= 0:
        return None
    s, si, c = (100.0 * v / total for v in vals)
    if si + 1.5 * c < 15:
        return "sand"
    if si + 2 * c < 30:
        return "loamy sand"
    if (7 <= c < 20 and s > 52) or (c < 7 and si < 50):
        return "sandy loam"
    if 7 <= c < 27 and 28 <= si < 50 and s <= 52:
        return "loam"
    if (si >= 50 and 12 <= c < 27) or (50 <= si < 80 and c < 12):
        return "silt loam"
    if si >= 80 and c < 12:
        return "silt"
    if 20 <= c < 35 and si < 28 and s > 45:
        return "sandy clay loam"
    if 27 <= c < 40 and 20 < s <= 45:
        return "clay loam"
    if 27 <= c < 40 and s <= 20:
        return "silty clay loam"
    if c >= 35 and s > 45:
        return "sandy clay"
    if c >= 40 and si >= 40:
        return "silty clay"
    if c >= 40:
        return "clay"
    return None


def aggregate_profile(horizons: list[dict], top: float = 0.0, bottom: float = 30.0) -> dict:
    """Medie pesate sullo spessore degli orizzonti dentro [top, bottom] (cm).

    `horizons`: dict con chiavi top, bottom, sand, silt, clay, caco3, ph
    (valori numerici o None). Ogni variabile usa solo gli orizzonti in cui c'è.
    """

    def weight(h) -> float:
        t, b = h.get("top"), h.get("bottom")
        if t is None or b is None or b <= t:
            return 0.0
        return max(0.0, min(b, bottom) - max(t, top))

    out: dict = {"n": 0, "depth_cm": 0.0}
    acc = defaultdict(float)
    wsum = defaultdict(float)
    tex_w = 0.0
    tex = {"sand": 0.0, "silt": 0.0, "clay": 0.0}
    covered = 0.0
    for h in horizons:
        w = weight(h)
        if w <= 0:
            continue
        out["n"] += 1
        covered += w
        for k in ("ph", "caco3"):
            if h.get(k) is not None:
                acc[k] += w * h[k]
                wsum[k] += w
        if all(h.get(k) is not None for k in ("sand", "silt", "clay")):
            for k in tex:
                tex[k] += w * h[k]
            tex_w += w
    out["depth_cm"] = round(covered, 1)
    for k in ("ph", "caco3"):
        out[k] = round(acc[k] / wsum[k], 2) if wsum[k] > 0 else None
    if tex_w > 0:
        for k in tex:
            out[k] = round(tex[k] / tex_w, 1)
        out["usda"] = usda_texture(out["sand"], out["silt"], out["clay"])
    else:
        out.update(sand=None, silt=None, clay=None, usda=None)
    return out


def dominant_soils(rows: list[dict], cfg: dict) -> dict[str, dict]:
    """Delineazione -> riga del suolo con percentuale più alta (a parità vince la prima)."""
    best: dict[str, dict] = {}
    for r in rows:
        k = norm_key(r.get(cfg["soils_join"]))
        if k is None:
            continue
        perc = _num(r.get(cfg["soil_perc"])) or 0.0
        if k not in best or perc > (_num(best[k].get(cfg["soil_perc"])) or 0.0):
            best[k] = r
    return best


def site_horizons(rows: list[dict], cfg: dict) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        k = norm_key(r.get(cfg["site_id"]))
        if k is None:
            continue
        out[k].append(
            {
                name: _num(r.get(cfg[name]))
                for name in ("top", "bottom", "sand", "silt", "clay", "caco3", "ph")
            }
        )
    return out


# --- lettura file -------------------------------------------------------------


def _read_sheet(path: Path, sheet: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ConfigError(f"{path}: foglio '{sheet}' non trovato (fogli: {', '.join(wb.sheetnames)})")
    it = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = [dict(zip(header, r)) for r in it]
    wb.close()
    return rows


def _check_columns(rows: list[dict], cols: list[str], where: str) -> None:
    have = set(rows[0]) if rows else set()
    missing = [c for c in cols if c not in have]
    if missing:
        raise ConfigError(f"{where}: colonne {missing} non trovate (presenti: {', '.join(sorted(have))})")


def _find_one(src: Source, key: str, pattern: str) -> Path:
    from ..paths import resolve

    v = src.prep(key)
    if v:
        p = resolve(v, src.base_dir)
        if not p.exists():
            raise ConfigError(f"[{src.id}] prepare.{key} non esiste: {p}")
        return p
    found = [p for p in src.cache_dir.rglob(pattern) if p.name != OUT_NAME]
    if len(found) != 1:
        raise NotConfigured(
            f"[{src.id}] per prepare serve un solo file {pattern} in {src.cache_dir} (trovati {len(found)}): "
            f"scarica il pacchetto (fetch {src.id}) o indica prepare.{key} in sources.toml."
        )
    return found[0]


def prepare_suoli(src: Source, *, log=print) -> Path:
    """Scrive data/raw/<id>/suoli_0_30.gpkg con le proprietà 0-30 cm per delineazione."""
    import geopandas as gpd
    import pyogrio

    from ..fetch.download import write_manifest

    miss = src.missing("prepare", *REQUIRED)
    if miss:
        raise NotConfigured(f"[{src.id}] per prepare mancano in sources.toml: {', '.join(miss)}")
    cfg = {k: src.prep(k) for k in REQUIRED}
    depth = float(src.prep("depth_cm", 30))
    shp = _find_one(src, "raw_shp", "*.shp")
    xlsx = _find_one(src, "raw_xlsx", "*.xlsx")
    log(f"[{src.id}] carta: {shp}\n[{src.id}] profili: {xlsx}")

    fields = [str(f) for f in pyogrio.read_info(shp)["fields"]]
    if cfg["shp_join"] not in fields:
        raise ConfigError(f"{shp}: attributo '{cfg['shp_join']}' non trovato (presenti: {', '.join(fields)})")
    soils = _read_sheet(xlsx, cfg["soils_sheet"])
    _check_columns(
        soils,
        [cfg[k] for k in ("soils_join", "soil_code", "soil_name", "soil_perc", "soil_site")],
        f"{xlsx}:{cfg['soils_sheet']}",
    )
    sites = _read_sheet(xlsx, cfg["sites_sheet"])
    _check_columns(
        sites,
        [cfg[k] for k in ("site_id", "top", "bottom", "sand", "silt", "clay", "caco3", "ph")],
        f"{xlsx}:{cfg['sites_sheet']}",
    )
    dom = dominant_soils(soils, cfg)
    hor = site_horizons(sites, cfg)

    gdf = gpd.read_file(shp)
    new = defaultdict(list)
    for v in gdf[cfg["shp_join"]]:
        d = dom.get(norm_key(v))
        site = norm_key(d.get(cfg["soil_site"])) if d else None
        agg = aggregate_profile(hor.get(site, []), 0.0, depth) if site else {}
        new["suolo_prev"].append(norm_key(d.get(cfg["soil_code"])) if d else None)
        new["nome_suolo_prev"].append(norm_key(d.get(cfg["soil_name"])) if d else None)
        new["perc_prev"].append(_num(d.get(cfg["soil_perc"])) if d else None)
        new["id_sito"].append(site)
        new["n_oriz"].append(agg.get("n", 0))
        new["prof_cm"].append(agg.get("depth_cm"))
        new["ph_0_30"].append(agg.get("ph"))
        new["caco3_0_30"].append(agg.get("caco3"))
        new["sabbia_0_30"].append(agg.get("sand"))
        new["limo_0_30"].append(agg.get("silt"))
        new["argilla_0_30"].append(agg.get("clay"))
        new["usda_0_30"].append(agg.get("usda"))
    for k, vals in new.items():
        gdf[k] = vals
    for k in ("perc_prev", "prof_cm", "ph_0_30", "caco3_0_30", "sabbia_0_30", "limo_0_30", "argilla_0_30"):
        gdf[k] = gdf[k].astype("float64")

    out = src.cache_dir / OUT_NAME
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()
    gdf.to_file(out, driver="GPKG", layer=OUT_LAYER)
    n = len(gdf)
    stats = {k: int(gdf[k].notna().sum()) for k in ("ph_0_30", "caco3_0_30", "usda_0_30")}
    log(f"[{src.id}] {n} delineazioni -> {out}; con valore: {stats}")
    prev = src.manifest() or {}
    write_manifest(
        src,
        out,
        method="prepare",
        prepared_from={"shp": str(shp), "xlsx": str(xlsx), "depth_cm": depth},
        fetch=prev if prev.get("method") != "prepare" else prev.get("fetch"),
        counts={"delineazioni": n, **stats},
    )
    return out
